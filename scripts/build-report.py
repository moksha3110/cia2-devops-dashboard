"""Build an eight-page PDF and editable DOCX from recorded deployment evidence.

Usage: python scripts/build-report.py
Dependencies (report tooling only): reportlab, pillow, python-docx.
Run after reports/evidence/summary.json and the evidence images exist.
"""
import json
from pathlib import Path
from html import escape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, Table, TableStyle
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Polygon
from docx import Document
from docx.shared import Inches, Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image as PILImage, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
E = OUT / 'evidence'
data = json.loads((E / 'summary.json').read_text(encoding='utf-8'))
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyReport', fontName='Helvetica', fontSize=10, leading=14, spaceAfter=9))
styles.add(ParagraphStyle(name='CaptionReport', fontSize=8, leading=11, textColor=colors.HexColor('#475569'), spaceAfter=9))
styles['Title'].fontSize = 23
styles['Title'].leading = 28
styles['Title'].textColor = colors.HexColor('#134e4a')
styles['Title'].spaceAfter = 16
styles['Heading1'].textColor = colors.HexColor('#0f766e')
styles['Heading1'].spaceBefore = 8
styles['Heading1'].spaceAfter = 12
story = []
doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.header.paragraphs[0].text = 'DEVOPS CIA 2 | Moksha | 23011102074'
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15
for style_name in ['Heading 1', 'Heading 2']:
    doc.styles[style_name].paragraph_format.keep_with_next = True
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_p.add_run('Moksha | IoT-B, 4th Year | Page ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer_p._p.append(field)
doc.core_properties.author = 'Moksha'
doc.core_properties.subject = 'DevOps CIA 2 | Roll number 23011102074 | IoT-B | 4th Year'
markdown = []

def p(text):
    story.append(Paragraph(escape(text), styles['BodyReport']))
    doc.add_paragraph(text)
    markdown.append(text + '\n')

def h(text):
    story.append(Paragraph(escape(text), styles['Heading1']))
    doc.add_heading(text, level=1)
    markdown.append('\n## ' + text + '\n')

def image(name, caption, max_height=240):
    path = E / name
    if not path.exists():
        raise FileNotFoundError(f'Required evidence missing: {path}')
    with PILImage.open(path) as im:
        width, height = im.size
    scale = min(480 / width, max_height / height)
    story.append(Image(str(path), width=width*scale, height=height*scale))
    story.append(Paragraph(escape(caption), styles['CaptionReport']))
    doc.add_picture(str(path), width=Inches(min(6.3, width*scale/72)))
    doc.add_paragraph(caption, style='Caption')
    markdown.append(f'![{caption}](evidence/{name})\n')

def table(rows, widths=None):
    cells = [[Paragraph(escape(str(c)), styles['CaptionReport']) for c in row] for row in rows]
    t = Table(cells, colWidths=widths, hAlign='LEFT', repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e2e8f0')),
        ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
        ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
    story.extend([t, Spacer(1,10)])
    dt = doc.add_table(rows=0, cols=len(rows[0]))
    dt.style = 'Table Grid'
    for row in rows:
        for cell, text in zip(dt.add_row().cells, row):
            cell.text = str(text)
    for cell in dt.rows[0].cells:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'E2E8F0')
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    markdown.append('\n| ' + ' | '.join(map(str,rows[0])) + ' |')
    markdown.append('| ' + ' | '.join('---' for _ in rows[0]) + ' |')
    markdown.extend(['| ' + ' | '.join(map(str,row)) + ' |' for row in rows[1:]])
    markdown.append('')

def nextpage():
    story.append(PageBreak())
    doc.add_page_break()

# Page 1: problem, scope and architecture.
story.append(Paragraph('Design and Implementation of an End-to-End Cloud-Based CI/CD Pipeline', styles['Title']))
doc.add_heading('Design and Implementation of an End-to-End Cloud-Based CI/CD Pipeline', 0)
markdown.append('# Design and Implementation of an End-to-End Cloud-Based CI/CD Pipeline\n')
table([['DEVOPS CIA 2', 'STUDENT DETAILS'],
       ['Name', 'Moksha'], ['Roll number', '23011102074'],
       ['Class / section', 'IoT-B'], ['Year', '4th Year']], [145,335])
p('Technical report | Evidence date: 12 October 2026 (IST)')
h('1. Problem definition and architecture')
p('Manual compilation, testing, image creation and deployment are slow, inconsistent and difficult to reproduce. This implementation automates those operations for a Java web application: a push to main triggers validation, publishes a container image and updates a cloud-hosted Kubernetes Deployment. A successful run must prove both rollout readiness and the public application identity.')
p('The application is a lightweight DevOps Deployment Dashboard. It displays version, environment, deployment message and Git commit SHA, making the change from v1 to v2 visible. AWS was selected as the permitted cloud provider; Amazon EKS hosts Kubernetes, while GitHub Actions orchestrates delivery and GHCR stores images.')
d = Drawing(480,135)
labels = ['Git push', 'Maven + tests', 'Docker / GHCR', 'AWS EKS', 'Browser']
for i, label in enumerate(labels):
    x = i*97
    d.add(Rect(x,70,88,40,rx=5,fillColor=colors.HexColor('#e0f2f1'),strokeColor=colors.HexColor('#0f766e')))
    d.add(String(x+44,87,label,textAnchor='middle',fontSize=9))
    if i < 4:
        d.add(Line(x+88,90,x+96,90,strokeColor=colors.HexColor('#0f766e')))
d.add(String(240,42,'GitHub Actions: CI -> publish -> deploy -> verify',textAnchor='middle',fontSize=10))
d.add(String(240,20,'Short-lived AWS OIDC | SHA tag + immutable digest | Two replicas',textAnchor='middle',fontSize=9))
story.extend([d, Spacer(1,8)])
diagram = PILImage.new('RGB', (1600, 360), '#ffffff')
draw = ImageDraw.Draw(diagram)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 25)
for i, label in enumerate(labels):
    x = 15+i*320
    draw.rounded_rectangle((x,80,x+280,180), radius=12, fill='#e0f2f1', outline='#0f766e', width=3)
    draw.text((x+140,130), label, font=font, anchor='mm', fill='#134e4a')
    if i<4:
        draw.line((x+280,130,x+315,130), fill='#0f766e', width=4)
        draw.polygon([(x+315,130),(x+304,123),(x+304,137)], fill='#0f766e')
draw.text((800,235),'GitHub Actions: CI -> publish -> deploy -> verify',font=font,anchor='mm',fill='#134e4a')
draw.text((800,295),'OIDC authentication | SHA + digest | Two replicas | Public NLB',font=font,anchor='mm',fill='#134e4a')
diagram.save(E/'architecture.png')
doc.add_picture(str(E/'architecture.png'), width=Inches(6.3))
doc.add_paragraph('Architecture diagram: delivery flow and security controls.', style='Caption')
markdown.append('![Architecture diagram](evidence/architecture.png)\n')
p('The deployment uses two replicas with rolling updates. The LoadBalancer Service maps public port 80 to application port 8080. Health probes control traffic readiness, and the verifier checks that the public endpoint reports the exact commit deployed by the workflow.')

nextpage()
h('2. Tools selected and application implementation')
table([['Tool', 'Purpose and justification'],
 ['Java 21 / Spring Boot 3.5.16','LTS Java runtime, concise REST APIs, embedded server and Actuator health probes.'],
 ['Maven Wrapper 3.9.16 / JUnit 5','Repeatable builds and automated endpoint/application tests.'],
 ['Git / GitHub Actions','Version history and event-driven CI/CD with separate dependent jobs.'],
 ['Docker / GHCR','Portable container packaging and SHA-tagged registry publication.'],
 ['AWS EKS Auto Mode / Kubernetes','Managed Kubernetes with automatic compute and NLB integration.'],
 ['Ansible','Optional read-only deployment verification using kubectl and HTTP checks.']], [145,335])
p('GET / serves responsive static HTML/CSS/JavaScript. GET /api/info returns name, version, environment, message and optional commitSha. GET /actuator/health returns overall health; liveness and readiness endpoints support Kubernetes probes. Metadata is supplied by APP_NAME, APP_VERSION, APP_ENVIRONMENT, DEPLOYMENT_MESSAGE and COMMIT_SHA, keeping the application independent of its deployment environment.')
p('Eight automated tests cover context startup, property binding, optional SHA serialization, valid dashboard HTML, metadata JSON, health UP, liveness and readiness. MockMvc handles REST checks; an actual random-port HTTP request verifies the Spring welcome page. Tests set explicit metadata and require no database or external service.')
p('The multi-stage Dockerfile builds and tests source with a JDK, then copies only the executable JAR into a smaller JRE image. Runtime UID/GID 10001 is non-root. Kubernetes drops all Linux capabilities, disables privilege escalation, mounts the root filesystem read-only and provides a bounded writable /tmp volume.')
table([['Validation', 'Recorded result'], ['Windows Maven clean test / package','8 tests, 0 failures, 0 errors, 0 skipped; both commands passed.'], ['GitHub CI test reports',data['test_result']], ['Configuration checks','actionlint passed; five Kubernetes resources passed schema validation; Bash and Ansible syntax checks passed.']], [175,305])

nextpage()
h('3. CI/CD pipeline stages')
p('Pull requests to main run CI only. Pushes to main and manual dispatch run CI, publish, and (when DEPLOY_ENABLED=true) deployment. Job dependencies prevent a failing build from publishing an image or deploying. Repository variables hold identifiers; secrets and cloud access keys are absent from source.')
table([['Stage','Implementation'], ['Checkout / compile / test / package','Java 21 setup, Maven cache, ./mvnw clean verify; JUnit reports and executable JAR uploaded.'], ['Build and publish','Buildx builds linux/amd64; GHCR login uses GITHUB_TOKEN with packages:write. Image tag is the full commit SHA.'], ['Authenticate / connect','GitHub OIDC assumes a role allowed to describe only the target EKS cluster; kubectl matches its minor version.'], ['Deploy / verify','Render image@sha256:digest and COMMIT_SHA; apply manifests; verify two ready updated replicas, pods and public HTTP metadata.'], ['Evidence / Ansible','Save cluster/pod/service/HTTP evidence; manual runs also execute the Ansible verification playbook.']], [130,350])
image('pipeline-v2.png', 'Figure 1. Actual GitHub Actions page for the automatic v2 run.', 275)
p('Workflow source: .github/workflows/ci-cd.yml. The same published digest is deployed, so retagging an image cannot silently change the running artifact. The dashboard delivery path is explanatory; the GitHub run is the authoritative pipeline record.')

nextpage()
h('4. Cloud resources and deployment approach')
p('CloudFormation stack cia2-eks created a dedicated VPC, two public subnets in different Mumbai availability zones, an internet gateway, infrastructure IAM roles and an EKS Auto Mode cluster. Kubernetes version 1.36 was in standard support at creation. Public worker IPs provide egress to GHCR without a NAT gateway; security groups control inbound access.')
image('cloud-kubernetes.png', 'Figure 2. Browser-rendered evidence view of actual AWS CLI and workflow kubectl output; this is not an AWS Console screenshot.', 320)
p('The Deployment requests 100m CPU and 256Mi memory per replica, with limits of 500m CPU and 512Mi. RollingUpdate uses maxSurge=1 and maxUnavailable=0. A startup probe permits warm-up, readiness removes unhealthy pods from traffic and liveness detects failed processes. The Service requests an internet-facing Auto Mode network load balancer.')
p('The namespace and RoleBinding were bootstrapped once. A temporary cluster-admin access policy was explicitly approved, used for bootstrap and removed afterward; its temporary workflow was also deleted. Routine CI/CD uses namespace-scoped RBAC rather than cluster-admin. The local root session could provision AWS resources but could not assume the Kubernetes administrator role.')
p('Public application endpoint: ' + data['url'])

nextpage()
h('5. Initial v1 deployment and registry evidence')
image('dashboard-v1.png', 'Figure 3. Actual live v1 dashboard after the first verified deployment.', 265)
image('registry.png', 'Figure 4. Actual public GHCR package page showing published container versions.', 230)
p('Initial verified commit: ' + data['v1_sha'])
p('Manual verification run: ' + data['v1_run_url'])
p('The registry package is public, so the cluster can pull it without storing a personal access token. The application verifier required HTTP 200, HTML content, health UP, aws-eks environment, expected version and the expected commit SHA. Optional Ansible verification result: ' + data['ansible_result'])

nextpage()
h('6. Code change and automatic redeployment')
p('After capturing v1, the application HTML subtitle was changed to describe automatic GitHub-to-AWS delivery. APP_VERSION changed from v1 to v2, and DEPLOYMENT_MESSAGE identified the automatic redeployment. These source/configuration changes were committed and pushed to main. No manual kubectl deployment command was used for the update; the push triggered CI, image publication and rollout.')
image('dashboard-v2-card.png', 'Figure 5. Actual live v2 dashboard detail with updated version, message and commit SHA.', 290)
table([['Traceability item','Recorded value'], ['v1 commit',data['v1_sha']], ['v2 source change','ff59e4b8600af134c1f1c20a47370d6045e56efe'], ['v2 release commit',data['v2_sha']], ['Automatic v2 workflow',data['v2_run_url']], ['v2 outcome',data['v2_result']]], [120,360])
p('The source-change push first succeeded in run 38229922992. The recorded final release also includes a follow-up Service annotation enabling cross-zone NLB routing; it was deployed automatically by the linked push run.')
p('The recorded image reference includes both the full commit SHA tag and registry digest. Kubernetes evidence records two updated, ready, available replicas. Public HTTP verification compares COMMIT_SHA against the workflow commit, proving the visible version came from the expected source revision rather than a stale pod or cached page.')

nextpage()
h('7. Challenges, security and limitations')
table([['Challenge','Resolution'], ['First Maven download failed on runner','The configured Maven URL was reachable; a retry passed. No test success was claimed for the failed attempt.'], ['Root could not assume administrator role','Used explicitly approved one-time GitHub OIDC bootstrap; removed cluster-admin policy immediately afterward.'], ['GHCR package visibility','Confirmed public package access before pulling from EKS; avoided a persistent registry password.'], ['New GitHub OIDC subject format','Verified immutable owner/repository IDs through the GitHub API and restricted trust to the exact main-branch subject.']], [170,310])
p('Workflow permissions are scoped per job: CI reads source, publication writes packages, and deployment requests OIDC tokens. Untrusted pull requests do not publish or deploy. The AWS deployment role has only eks:DescribeCluster on the target cluster; Kubernetes Role/RoleBinding limits application operations to devops-dashboard. Containers run as non-root with restricted filesystem and capabilities.')
p('Limitations: the demonstration uses public HTTP without TLS, a public Kubernetes API protected by authentication, and public worker subnets. Production should add HTTPS, a domain, stronger branch controls, private networking/runner access where appropriate, image vulnerability scanning, monitoring and alerting. Action versions and base image tags are versioned but not all pinned by digest. Successful rollout does not prove fault tolerance or load performance; no load test or failure-injection test was performed.')
p('EKS, EC2, EBS, public IPv4 and the NLB incur charges while retained. This implementation does not claim free-tier eligibility. After assessment, delete the Service and wait for its load balancer to disappear, remove workloads/compute and then delete the dedicated CloudFormation stack; inspect for residual resources. Cleanup requires an intentional operation and has not been performed.')

nextpage()
h('8. Submission evidence, live demonstration and references')
p('Submission includes the public source repository, Maven Wrapper, application/tests, Dockerfile, .dockerignore, GitHub Actions workflow, Kubernetes manifests, optional Ansible playbook/inventory, CloudFormation template and detailed setup/troubleshooting guides. The reports/evidence directory contains original screenshots, recorded API/CLI evidence and test results used by this report.')
p('Live demonstration sequence: open the dashboard and note the current version/SHA; open GitHub Actions; edit APP_VERSION and DEPLOYMENT_MESSAGE in k8s/deployment.yaml; commit and push main; watch CI, publish and deploy complete; refresh the public dashboard and compare its SHA with the commit. Explain the workflow guards, image digest, readiness probes, RBAC and rolling-update settings during the viva.')
table([['Deliverable','Evidence location'], ['Source and configuration','https://github.com/moksha3110/cia2-devops-dashboard'], ['Pipeline execution',data['v2_run_url']], ['Working application',data['url']], ['Tests / cloud state / screenshots','reports/evidence/ and GitHub run artifacts'], ['Setup / troubleshooting / viva','docs/deployment-guide.md, troubleshooting.md, architecture.md, evidence-checklist.md']], [145,335])
p('Evidence integrity: dashboard, repository, pipeline and registry pictures are actual browser captures. Figure 2 explicitly presents saved AWS/kubectl output in an evidence view; it does not impersonate the AWS Console. No AI-generated image or invented deployment result is used. Raw workflow artifacts remain available subject to the configured retention period. Optional instructor-specific Console screenshots may be added separately.')
h('Primary references')
for label, url in [
 ('AWS EKS Auto Mode CLI setup','https://docs.aws.amazon.com/eks/latest/userguide/automode-get-started-cli.html'),
 ('AWS EKS access entries','https://docs.aws.amazon.com/eks/latest/userguide/access-entries.html'),
 ('GitHub OIDC with AWS','https://docs.github.com/en/actions/how-tos/secure-your-work/security-harden-deployments/oidc-in-aws'),
 ('GitHub Container Registry','https://docs.github.com/en/packages/working-with-a-github-packages-registry/working-with-the-container-registry'),
 ('Docker multi-stage builds','https://docs.docker.com/build/building/multi-stage/'),
 ('EKS pricing','https://aws.amazon.com/eks/pricing/')]:
    story.append(Paragraph(f'{escape(label)}: <link href="{url}">{escape(url)}</link>', styles['CaptionReport']))
    doc.add_paragraph(label + ': ' + url)
    markdown.append(label + ': ' + url + '\n')

def footer(canvas, document):
    canvas.setFont('Helvetica',8)
    canvas.setFillColor(colors.HexColor('#64748b'))
    canvas.setStrokeColor(colors.HexColor('#cbd5e1'))
    canvas.line(42,38,A4[0]-42,38)
    canvas.drawString(42,25,'Moksha | 23011102074 | IoT-B, 4th Year | DevOps CIA 2')
    canvas.drawRightString(A4[0]-42,25,f'Page {document.page}')

SimpleDocTemplate(str(OUT/'CIA2-Technical-Report.pdf'),pagesize=A4,
    rightMargin=42,leftMargin=42,topMargin=38,bottomMargin=42,
    title='DevOps CIA 2 - End-to-End Cloud-Based CI/CD Pipeline',author='Moksha').build(story,onFirstPage=footer,onLaterPages=footer)
doc.save(OUT/'CIA2-Technical-Report.docx')
(OUT/'CIA2-Technical-Report.md').write_text('\n'.join(markdown),encoding='utf-8')
print('Generated PDF, editable DOCX and Markdown report in reports/.')
